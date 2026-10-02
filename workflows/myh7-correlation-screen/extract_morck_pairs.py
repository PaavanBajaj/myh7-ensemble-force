"""Extract only the selected paired velocity numbers from Morck Figure 3 source data.

Requires openpyxl and the checksum-pinned publisher workbook. The workbook is
not bundled here; this script writes a compact measurement/provenance table.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from datetime import datetime
from pathlib import Path
from statistics import mean, stdev

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


EXPECTED_SHA256 = "01020dd880b41691d1a8b7edce2e15125ad504d613e593450a693fd59f4de4f7"
SHEET = "Fig. 3F In vitro motility"
EXPECTED_COUNTS = {"D778V": 5, "L781P": 5, "S782N": 9, "A797T": 7, "F834L": 6}
PUBLISHER_URL = "https://cdn.elifesciences.org/articles/76805/elife-76805-fig3-data1-v1.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    observed = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    if observed != EXPECTED_SHA256:
        raise ValueError(f"Morck source workbook SHA-256 mismatch: {observed}")
    sheet = load_workbook(args.workbook, read_only=True, data_only=True)[SHEET]
    source_rows = []
    date = slide = ""
    for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
        label = cells[8] if len(cells) > 8 else None
        if isinstance(label, datetime):
            date = label.date().isoformat()
        elif isinstance(label, str) and label.startswith("Slide"):
            slide = label.replace(" ", "").lower()
        ratio = cells[14] if len(cells) > 14 else None
        if isinstance(ratio, (int, float)):
            source_rows.append((number, date, slide, label, cells[9], float(ratio)))
    results = []
    for column, (variant, expected_n) in enumerate(EXPECTED_COUNTS.items(), 2):
        values = []
        for summary_row in range(5, 14):
            rounded = sheet.cell(summary_row, column).value
            if not isinstance(rounded, (int, float)):
                continue
            matches = [row for row in source_rows if row[3] == variant and abs(row[5] - rounded) < 1e-7]
            if len(matches) != 1:
                raise ValueError(f"{variant} {get_column_letter(column)}{summary_row} lacks a unique absolute pair")
            mutant_row, date, slide, _, mutant, ratio = matches[0]
            wt_row = mutant_row - 1
            while wt_row > 0:
                candidate = sheet.cell(wt_row, 9).value
                if isinstance(candidate, str) and "WT" in candidate:
                    break
                wt_row -= 1
            if wt_row < 1:
                raise ValueError(f"{variant} workbook row {mutant_row} lacks a WT pair")
            wt = sheet.cell(wt_row, 10).value
            if not isinstance(wt, (int, float)) or not isinstance(mutant, (int, float)):
                raise ValueError(f"{variant} workbook row {mutant_row} lacks numeric absolutes")
            exact = float(mutant) / float(wt)
            if abs(exact - ratio) > 1e-12:
                raise ValueError(f"{variant} workbook row {mutant_row} ratio/absolute mismatch")
            values.append(exact)
            results.append({
                "measurement_id": f"VM{len(results) + 1:03d}",
                "variant": variant,
                "source_id": "morck2022",
                "source_workbook": PUBLISHER_URL,
                "source_sheet": SHEET,
                "summary_cell": f"{get_column_letter(column)}{summary_row}",
                "wt_workbook_row": wt_row,
                "mutant_workbook_row": mutant_row,
                "date": date,
                "slide_id": f"{date}-{slide}",
                "wt_context_id": f"morck2022-fig3f-row{wt_row}",
                "wt_velocity_nm_s": int(wt),
                "mutant_velocity_nm_s": int(mutant),
                "v_rel": f"{exact:.12f}",
                "biological_prep_id": "unavailable",
            })
        if len(values) != expected_n:
            raise ValueError(f"{variant} has {len(values)} selected pairs, expected {expected_n}")
        print(f"{variant}: {len(values)} selected slide/channel ratios, mean={mean(values):.9f}, sample_SD={stdev(values):.9f}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(results[0]))
        writer.writeheader()
        writer.writerows(results)


if __name__ == "__main__":
    main()
